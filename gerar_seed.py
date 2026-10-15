# -*- coding: utf-8 -*-
"""Gera supabase/seed.sql a partir do Excel "Acompanhamento Consultores - Fertilizer AVC - LT Facil.xlsx".

Uso:  python3 gerar_seed.py "caminho/Acompanhamento Consultores - Fertilizer AVC - LT Facil.xlsx" > ../supabase/seed.sql
Rode o seed.sql no SQL Editor do Supabase depois do schema.sql. Pode ser rodado de novo: ele substitui os dados.
"""
import sys, datetime as dt
from openpyxl import load_workbook

src = sys.argv[1]
wb = load_workbook(src, data_only=True)

def s(v):
    if v is None or v == "":
        return "null"
    if isinstance(v, (dt.datetime, dt.date)):
        return "'%s'" % v.strftime("%Y-%m-%d")
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, (int, float)):
        return repr(v)
    return "'" + str(v).replace("'", "''") + "'"

out = ["-- seed gerado em %s a partir de %s" % (dt.date.today().isoformat(), src.split("/")[-1]), "begin;"]

# parametros
pa = wb["Parâmetros"]
out.append("delete from public.parametros;")
for chave, ref in (("data_referencia", "B4"), ("inicio_testes", "B5"), ("cliente", "B6"), ("projeto", "B7"), ("gerente", "B8"), ("diretor", "B9")):
    v = pa[ref].value
    out.append("insert into public.parametros (chave, valor) values (%s, %s);" % (s(chave), s(v.strftime("%Y-%m-%d") if isinstance(v, (dt.date, dt.datetime)) else v)))
fer = [pa.cell(r, 4).value for r in range(4, 24) if isinstance(pa.cell(r, 4).value, (dt.date, dt.datetime))]
out.append("insert into public.parametros (chave, valor) values ('feriados', %s);" % s(";".join(f.strftime("%Y-%m-%d") for f in fer)))

# equipe
eq = wb["Equipe"]
out.append("delete from public.equipe;")
for i, r in enumerate(range(5, 17)):
    if not eq.cell(r, 1).value:
        continue
    vals = [eq.cell(r, c).value for c in range(1, 8)]
    out.append("insert into public.equipe (ordem, frente, key_user, area_key_user, consultor, apoio, contato, obs) values (%d, %s);" % (i, ", ".join(s(v) for v in vals)))
out.append("delete from public.papeis;")
for i, r in enumerate(range(20, 25)):
    if not eq.cell(r, 1).value:
        continue
    out.append("insert into public.papeis (ordem, papel, nome, responsabilidade) values (%d, %s, %s, %s);" % (i, s(eq.cell(r, 1).value), s(eq.cell(r, 2).value), s(eq.cell(r, 3).value)))

# atividades
out.append("delete from public.atividades;")
def atividades(sheet, origem):
    ws = wb[sheet]
    n = 0
    for r in range(6, ws.max_row + 1):
        bloco, codigo, atividade = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value
        if not codigo or not atividade or bloco == "Total":
            if bloco == "Total":
                break
            continue
        lado, consultor, key_user, status, pct = [ws.cell(r, c).value for c in (4, 5, 6, 7, 8)]
        inicio, previsao, conclusao = [ws.cell(r, c).value for c in (9, 10, 11)]
        precedentes, ponto, obs = ws.cell(r, 14).value, ws.cell(r, 16).value, ws.cell(r, 17).value
        pct = round(float(pct or 0) * 100, 2)
        out.append("insert into public.atividades (origem, ordem, bloco, codigo, atividade, lado, consultor, key_user, status, pct, inicio, previsao, conclusao, precedentes, ponto_atencao, obs) values (%s);"
                   % ", ".join([s(origem), str(n), s(bloco), s(str(codigo)), s(atividade), s(lado), s(consultor), s(key_user), s(status or "Não iniciado"), str(pct), s(inicio), s(previsao), s(conclusao), s(precedentes), s(ponto), s(obs)]))
        n += 1
    return n
nt = atividades("Cronograma Template", "template")
np_ = atividades("Cronograma Portal", "portal")
nc = atividades("Customizações", "custom")

# roteiro
rt = wb["Roteiro de Testes"]
out.append("delete from public.roteiro;")
n = 0
for r in range(6, rt.max_row + 1):
    cod, fluxo, passo = rt.cell(r, 1).value, rt.cell(r, 2).value, rt.cell(r, 4).value
    if not cod or not passo:
        continue
    vals = [rt.cell(r, c).value for c in range(2, 15)]
    fluxo, cenario, passo, sistema, pre, dados, esperado, executor, consultor, data, resultado, evid, obs = vals
    out.append("insert into public.roteiro (ordem, codigo, fluxo, cenario, passo, sistema, pre_condicao, dados_teste, resultado_esperado, executor, consultor, data_execucao, resultado, evidencia, obs) values (%s);"
               % ", ".join([str(n), s(cod), s(fluxo), s(cenario), s(passo), s(sistema), s(pre), s(dados), s(esperado), s(executor), s(consultor), s(data), s(resultado or "Não executado"), s(evid), s(obs)]))
    n += 1
nr = n

# pontos
pt = wb["Pontos de Atenção"]
out.append("delete from public.pontos;")
n = 0
for r in range(6, pt.max_row + 1):
    cod, desc = pt.cell(r, 1).value, pt.cell(r, 5).value
    if not cod or not desc:
        continue
    data, frente, tipo, desc, impacto, resp, prazo, status, prec, dep, bol, resol = [pt.cell(r, c).value for c in range(2, 14)]
    out.append("insert into public.pontos (codigo, data, frente, tipo, descricao, impacto, responsavel, prazo, status, precedente, dependentes, no_boletim, resolucao) values (%s);"
               % ", ".join([s(cod), s(data or dt.date.today()), s(frente), s(tipo), s(desc), s(impacto), s(resp), s(prazo), s(status or "Aberto"), s(prec), s(dep), "true" if (bol == "Sim") else "false", s(resol)]))
    n += 1
npt = n
out.append("commit;")
out.append("-- atividades: template %d, portal %d, custom %d | roteiro %d | pontos %d" % (nt, np_, nc, nr, npt))
print("\n".join(out))
